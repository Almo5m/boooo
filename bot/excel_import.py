import json
import io
from openpyxl import load_workbook
from database import get_session, Question


def import_questions_from_excel(file_bytes: bytes, subject_id: int) -> tuple[int, list[str]]:
    """
    شكل الإكسيل المتوقع (بالترتيب، بغض النظر عن نص العنوان في أول صف):
    A: نص السؤال | B: اختيار 1 | C: اختيار 2 | D: اختيار 3 (اختياري) | E: اختيار 4 (اختياري)
    F: رقم الإجابة الصحيحة (0 = الأول) | G: شرح (اختياري) | H: الصعوبة (اختياري: easy/medium/hard)

    بترجع (عدد الأسئلة اللي اتضافت, لستة أخطاء لو فيه صفوف اتجاهلت)
    """
    added = 0
    errors = []
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active

    session = get_session()
    try:
        for row_idx, row in enumerate(ws.iter_rows(min_row=1, values_only=True), start=1):
            if not row or not row[0]:
                continue
            question_text = str(row[0]).strip()

            # تجاهل صف العنوان لو أول خلية مش سؤال حقيقي والعمود السادس مش رقم
            correct_raw = row[5] if len(row) > 5 else None
            if row_idx == 1 and not _is_int(correct_raw):
                continue

            if not _is_int(correct_raw):
                errors.append(f"صف {row_idx}: عمود الإجابة الصحيحة (F) لازم يكون رقم")
                continue

            options = []
            for col in range(1, 5):  # B, C, D, E
                val = row[col] if len(row) > col else None
                if val is not None and str(val).strip():
                    options.append(str(val).strip())

            if len(options) < 2:
                errors.append(f"صف {row_idx}: لازم اختيارين على الأقل")
                continue

            correct_index = int(correct_raw)
            if correct_index >= len(options):
                errors.append(f"صف {row_idx}: رقم الإجابة الصحيحة أكبر من عدد الاختيارات")
                continue

            explanation = str(row[6]).strip() if len(row) > 6 and row[6] else ""
            difficulty = str(row[7]).strip().lower() if len(row) > 7 and row[7] else "medium"
            if difficulty not in ("easy", "medium", "hard"):
                difficulty = "medium"

            q = Question(
                subject_id=subject_id,
                text=question_text,
                options=json.dumps(options, ensure_ascii=False),
                correct_index=correct_index,
                explanation=explanation,
                difficulty=difficulty,
            )
            session.add(q)
            added += 1

        session.commit()
    finally:
        session.close()

    return added, errors


def _is_int(val) -> bool:
    if val is None:
        return False
    try:
        int(val)
        return True
    except (ValueError, TypeError):
        return False
